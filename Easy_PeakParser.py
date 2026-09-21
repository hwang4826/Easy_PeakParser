import os
import sys
import csv
import webbrowser
import ctypes
from datetime import datetime, timedelta
import tkinter as tk
from tkinter import filedialog, messagebox
from tkinterdnd2 import TkinterDnD, DND_FILES
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ================================
# 작업 표시줄 아이콘 고정 버그 해결 (Windows 전용)
# ================================
try:
    myappid = 'tuk.mecha.easy_peakparser.v1'
    ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(myappid)
except:
    pass

# ================================
# PyInstaller 임시 폴더 경로 추적 함수
# ================================
def resource_path(relative_path):
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)

mfc_files = []
peak_files = []

# ================================
# 시간 포맷 정밀 파싱 함수
# ================================
def parse_mfc_time(t_str):
    formats = [
        "%y%m%d %H:%M:%S.%f",
        "%y%m%d %H:%M:%S",
        "%Y-%m-%d %H:%M:%S.%f",
        "%Y-%m-%d %H:%M:%S"
    ]
    t_str = str(t_str).strip()
    for fmt in formats:
        try:
            return datetime.strptime(t_str, fmt)
        except ValueError:
            continue
    return None

def parse_peak_time(t_str):
    try:
        parts = str(t_str).strip().split(':')
        if len(parts) >= 7:
            ms = parts[6]
            if len(ms) == 3: ms += "000"
            t_str_fixed = ":".join(parts[:6]) + "." + ms
            return datetime.strptime(t_str_fixed, "%Y:%m:%d:%H:%M:%S.%f")
    except:
        pass
    return None

def get_formatted_filename(pattern, base_time=None):
    if base_time is None:
        base_time = datetime.now()
        
    res = pattern
    res = res.replace("yyyy", base_time.strftime("%Y")).replace("yy", base_time.strftime("%y"))
    res = res.replace("MM", base_time.strftime("%m")).replace("dd", base_time.strftime("%d"))
    res = res.replace("HH", base_time.strftime("%H")).replace("hh", base_time.strftime("%H"))
    res = res.replace("mm", base_time.strftime("%M")).replace("ss", base_time.strftime("%S"))
    for char in ['\\', '/', ':', '*', '?', '"', '<', '>', '|']:
        res = res.replace(char, "-" if char == ':' else "_")
    return res

# ================================
# MFC 데이터 추출기
# ================================
def extract_mfc_events(filepath):
    events = []
    ext = filepath.lower()
    
    if ext.endswith('.csv'):
        with open(filepath, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            rows = list(reader)
            if not rows: return events
            
            if 'BG_Gas' in reader.fieldnames and 'Prev_Flow' in reader.fieldnames:
                for row in rows:
                    dt = parse_mfc_time(row['Time'])
                    if dt:
                        events.append({
                            'time': dt, 'bg_gas': row['BG_Gas'], 'bg_flow': float(row['BG_Flow']),
                            'ctrl_gas': row['Ctrl_Gas'], 'prev_flow': float(row['Prev_Flow']), 'curr_flow': float(row['Curr_Flow'])
                        })
            elif 'Ar_SetPT' in reader.fieldnames: 
                ar_setpts = [float(row['Ar_SetPT']) for row in rows]
                n2_setpts = [float(row['N2_SetPT']) for row in rows]
                ar_is_bg = (min(ar_setpts) == max(ar_setpts))
                n2_is_bg = (min(n2_setpts) == max(n2_setpts))

                if ar_is_bg and not n2_is_bg:
                    bg_gas, ctrl_gas = 'Ar', 'N2'
                    bg_col, ctrl_col = 'Ar_SetPT', 'N2_SetPT'
                elif n2_is_bg and not ar_is_bg:
                    bg_gas, ctrl_gas = 'N2', 'Ar'
                    bg_col, ctrl_col = 'N2_SetPT', 'Ar_SetPT'
                else:
                    return events

                bg_flow = float(rows[0][bg_col])
                prev_ctrl = float(rows[0][ctrl_col])
                for i in range(1, len(rows)):
                    curr_ctrl = float(rows[i][ctrl_col])
                    if curr_ctrl != prev_ctrl:
                        dt = parse_mfc_time(rows[i-1]['Time'])
                        if dt:
                            events.append({
                                'time': dt, 'bg_gas': bg_gas, 'bg_flow': bg_flow,
                                'ctrl_gas': ctrl_gas, 'prev_flow': prev_ctrl, 'curr_flow': curr_ctrl
                            })
                        prev_ctrl = curr_ctrl

    elif ext.endswith('.xlsx'):
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active
        headers = [str(cell.value) for cell in ws[1]]
        if 'BG_Gas' in headers:
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[1]: continue
                t_val = row[1]
                dt = t_val if isinstance(t_val, datetime) else parse_mfc_time(str(t_val))
                if dt:
                    events.append({
                        'time': dt, 'bg_gas': str(row[2]), 'bg_flow': float(row[3]),
                        'ctrl_gas': str(row[4]), 'prev_flow': float(row[5]), 'curr_flow': float(row[6])
                    })
    return events

# ================================
# 핵심 메인 프로세스
# ================================
def process_data():
    if not mfc_files or not peak_files:
        messagebox.showerror("에러", "MFC 유량 제어 기록 파일과 파장 기록 파일이 모두 입력되어야 합니다.")
        return

    output_dir = filedialog.askdirectory(title="파장 데이터 변환 결과를 각각 저장할 '폴더'를 선택하세요")
    if not output_dir: return

    all_events = []
    for f in mfc_files:
        all_events.extend(extract_mfc_events(f))
    all_events.sort(key=lambda x: x['time'])

    success_count = 0
    fail_list = []

    for peak_file in peak_files:
        headers_info = []
        data_rows = []
        is_data = False

        with open(peak_file, 'r', encoding='utf-8') as f:
            for line in f:
                line = line.strip()
                if not line: continue
                if not is_data:
                    if line == "[Data]": is_data = True
                    elif line.startswith("Wavelength"):
                        try:
                            parts = line.split(',')
                            wl = parts[0].split('=')[1]
                            el = parts[1].split('=')[1]
                            headers_info.append(f"{wl} / {el}")
                        except: pass
                else:
                    if not line.startswith("Timestamp"):
                        parts = line.split(',')
                        if len(parts) >= 2:
                            dt = parse_peak_time(parts[0])
                            if dt:
                                clean_raw = parts[:2 + len(headers_info)]
                                data_rows.append({'dt': dt, 'raw': clean_raw})
        
        if not data_rows:
            fail_list.append(f"{os.path.basename(peak_file)} (데이터 없음)")
            continue

        start_dt = data_rows[0]['dt']
        end_dt = data_rows[-1]['dt']
        
        tolerance = timedelta(minutes=5)
        matched_events = [e for e in all_events if (start_dt - tolerance) <= e['time'] <= (end_dt + tolerance)]
        
        if not matched_events:
            fail_list.append(f"{os.path.basename(peak_file)} (시간대가 일치하는 MFC 기록 없음)")
            continue

        # 어설픈 사전 필터링 제거! 실제 시간대가 맞는 블록 내에서만 0 SCCM을 판단합니다.
        blocks_to_write = []
        used_red_indices = set()
        processed_events = []

        for e in matched_events:
            red_idx = -1
            for i, row in enumerate(data_rows):
                if row['dt'] <= e['time']:
                    red_idx = i
                else:
                    break
            
            if red_idx != -1:
                # 파장 데이터와 MFC 제어 시간 차이가 60초 이내인 정상 데이터만 처리
                if abs((e['time'] - data_rows[red_idx]['dt']).total_seconds()) > 60:
                    continue
                
                if red_idx not in used_red_indices:
                    used_red_indices.add(red_idx)
                    start_idx = max(0, red_idx - 10)
                    block = data_rows[start_idx : red_idx + 1]
                    
                    prev_val = int(e['prev_flow']) if e['prev_flow'].is_integer() else e['prev_flow']
                    label = f"{e['ctrl_gas']} {prev_val} SCCM"
                    
                    blocks_to_write.append({
                        'label': label,
                        'block': block
                    })
                    processed_events.append(e)

                    # [핵심 로직] 성공적으로 기록된 유효 이벤트의 셋포인트가 0 SCCM이라면, 여기서 한 사이클을 완벽히 종료
                    if prev_val == 0:
                        break

        if not blocks_to_write:
            fail_list.append(f"{os.path.basename(peak_file)} (매칭된 유효 데이터 없음)")
            continue

        # 최댓값과 가스 정보는 매칭에 성공한 첫 번째(시작점) 유효 데이터 기준으로 고정
        first_event = processed_events[0]
        bg_gas = first_event['bg_gas']
        bg_flow = first_event['bg_flow']
        ctrl_gas = first_event['ctrl_gas']
        
        max_ctrl = first_event['prev_flow']
        
        deltas = [abs(e['curr_flow'] - e['prev_flow']) for e in processed_events if abs(e['curr_flow'] - e['prev_flow']) > 0]
        delta_val = max(set(deltas), key=deltas.count) if deltas else 0
        delta_str = int(delta_val) if delta_val.is_integer() else delta_val
        
        bg_flow_str = int(bg_flow) if bg_flow.is_integer() else bg_flow
        max_ctrl_str = int(max_ctrl) if max_ctrl.is_integer() else max_ctrl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Peak Data"

        ws['B1'] = f"배경가스 {bg_gas} {bg_flow_str} sccm + 제어가스 {ctrl_gas} 0~{max_ctrl_str} sccm(Δ{delta_str} sccm)"
        ws['B1'].font = Font(size=15, bold=True)
        ws.merge_cells('B1:H1')
        ws['B1'].alignment = Alignment(horizontal="left", vertical="center")
        
        ws['B2'] = "[Data]"
        ws['D2'], ws['E2'] = "측정일 :", start_dt.strftime("%y%m%d")
        ws['F2'], ws['G2'] = "수정일 :", datetime.now().strftime("%y%m%d")
        
        for cell_ref in ['D2', 'E2', 'F2', 'G2']:
            ws[cell_ref].alignment = Alignment(horizontal="center", vertical="center")
            ws[cell_ref].font = Font(bold=True)

        ws.append(["", "Timestamp", "No"] + headers_info)

        palette_1 = {
            'normal': PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid"),
            'event': PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid"),
            'font_event': Font(color="FFFFFF", bold=True)
        }
        palette_2 = {
            'normal': PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
            'event': PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),
            'font_event': Font(color="000000", bold=True)
        }
        
        font_bold = Font(bold=True)
        center_align = Alignment(horizontal="center", vertical="center")

        col_palettes = {}
        col_palettes[1] = col_palettes[2] = col_palettes[3] = palette_1 
        
        current_pal_idx = 2 
        prev_el = None
        
        for i, header in enumerate(headers_info):
            col_idx = 4 + i
            el = header.split(' / ')[1].strip() if ' / ' in header else ""
            if prev_el is not None and el != prev_el:
                current_pal_idx = 1 if current_pal_idx == 2 else 2
            prev_el = el
            col_palettes[col_idx] = palette_1 if current_pal_idx == 1 else palette_2

        current_row = 4
        for b_dict in blocks_to_write:
            label = b_dict['label']
            block = b_dict['block']
            
            for i, b_row in enumerate(block):
                is_red = (i == len(block) - 1)
                row_label = "유량 변경" if is_red else (label if i == 0 else "")
                
                formatted_ts = b_row['dt'].strftime("%y%m%d %H:%M:%S.%f")[:-3]
                clean_raw = [formatted_ts, b_row['raw'][1]] + b_row['raw'][2:]
                
                row_data = [row_label] + clean_raw
                ws.append(row_data)

                for col_idx in range(1, len(row_data) + 1):
                    cell = ws.cell(row=current_row, column=col_idx)
                    cell.alignment = center_align
                    
                    pal = col_palettes.get(col_idx, palette_1)
                    
                    if is_red:
                        cell.fill = pal['event']
                        cell.font = pal['font_event']
                    else:
                        cell.fill = pal['normal']
                        if col_idx == 1: cell.font = font_bold
                current_row += 1

        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 21
        ws.column_dimensions['C'].width = 6
        for col_idx in range(4, 4 + len(headers_info)):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = 13
            
        ws.freeze_panes = 'D4' 

        raw_pattern = entry_filename.get().strip()
        if not raw_pattern:
            raw_pattern = "yyMMdd_HHmm"
            
        prefix = get_formatted_filename(raw_pattern, start_dt) if use_time_format_var.get() else raw_pattern
        
        if bg_gas == 'Ar':
            gas_suffix = f"Ar_B{bg_flow_str}_N2_C{max_ctrl_str}"
        else:
            gas_suffix = f"Ar_C{max_ctrl_str}_N2_B{bg_flow_str}"
            
        save_name = f"{prefix}_{gas_suffix}.xlsx"
        counter = 1
        while os.path.exists(os.path.join(output_dir, save_name)):
            save_name = f"{prefix}_{gas_suffix}_{counter}.xlsx"
            counter += 1
            
        wb.save(os.path.join(output_dir, save_name))
        success_count += 1

    msg = f"총 {len(peak_files)}개 중 {success_count}개 변환 성공!\n"
    if fail_list:
        msg += f"\n[ ❌ 실패 내역 ({len(fail_list)}개) ]\n" + "\n".join(fail_list)
        messagebox.showwarning("완료 (일부 실패 존재)", msg)
    else:
        messagebox.showinfo("완료", msg)

# ================================
# 폴더 드래그 앤 드롭 시 하위 CSV 파일까지 재귀 탐색하는 함수
# ================================
def process_dropped_paths(paths, target_list):
    for p in paths:
        if os.path.isdir(p):
            for root_dir, _, files in os.walk(p):
                for file in files:
                    if file.lower().endswith('.csv'):
                        full_path = os.path.join(root_dir, file).replace('\\', '/')
                        if full_path not in target_list:
                            target_list.append(full_path)
        elif os.path.isfile(p):
            if p.lower().endswith('.csv'):
                p = p.replace('\\', '/')
                if p not in target_list:
                    target_list.append(p)

# ================================
# GUI 리스트박스 관리 함수
# ================================
def update_lists():
    listbox_peak.delete(0, tk.END); listbox_mfc.delete(0, tk.END)
    for p in peak_files: listbox_peak.insert(tk.END, os.path.basename(p))
    for p in mfc_files: listbox_mfc.insert(tk.END, os.path.basename(p))

def drop_peak(event):
    process_dropped_paths(root.tk.splitlist(event.data), peak_files)
    update_lists()

def drop_mfc(event):
    process_dropped_paths(root.tk.splitlist(event.data), mfc_files)
    update_lists()

def browse_peak():
    for p in filedialog.askopenfilenames(title="Peak 파일 선택", filetypes=[("CSV", "*.csv")]):
        if p not in peak_files: peak_files.append(p)
    update_lists()

def browse_mfc():
    for p in filedialog.askopenfilenames(title="MFC 파일 선택", filetypes=[("CSV", "*.csv")]):
        if p not in mfc_files: mfc_files.append(p)
    update_lists()

def remove_peak():
    for i in reversed(listbox_peak.curselection()): del peak_files[i]
    update_lists()

def remove_mfc():
    for i in reversed(listbox_mfc.curselection()): del mfc_files[i]
    update_lists()

def clear_all():
    mfc_files.clear(); peak_files.clear()
    update_lists()

# ================================
# GitHub 링크 오픈 함수
# ================================
def open_github(event):
    webbrowser.open_new("https://github.com/hwang4826/Easy_PeakParser")

# ================================
# GUI 화면 구성 (TUK Blue Theme 적용 및 크기 5:5 고정)
# ================================
# 테마 색상표 정의
THEME_BG = "#F0F6FA"           
THEME_FRAME_BG = "#FFFFFF"     
THEME_ACCENT = "#005BAA"       
THEME_BTN = "#D4E6F1"          
THEME_BTN_ACTIVE = "#A9CCE3"   
THEME_ACTION = "#0078D7"       
THEME_ACTION_ACTIVE = "#005A9E" 

root = TkinterDnD.Tk()
# 버전 1.0.1 업데이트 반영
root.title("Easy_PeakParser v1.0.1")
root.geometry("720x580") 
root.iconbitmap(resource_path("Easy_PeakParser.ico"))
root.configure(bg=THEME_BG) 

# 파일 목록 프레임 창 크기 조절 시 동적 확장 적용
frame_lists = tk.Frame(root, bg=THEME_BG)
frame_lists.pack(pady=10, padx=10, fill="both", expand=True)

# 5:5 정확한 대칭을 위해 Grid 레이아웃 적용
frame_lists.columnconfigure(0, weight=1, uniform="equal_width")
frame_lists.columnconfigure(1, weight=1, uniform="equal_width")
frame_lists.rowconfigure(0, weight=1)

# [1. 파장 기록 영역 - 좌측 배치]
frame_peak = tk.LabelFrame(frame_lists, text="1. 파장 기록 파일 (*.csv)", bg=THEME_BG, fg=THEME_ACCENT, font=("", 10, "bold"), padx=5, pady=5)
frame_peak.grid(row=0, column=0, sticky="nsew", padx=5)

listbox_peak = tk.Listbox(frame_peak, selectmode=tk.EXTENDED, height=10, relief="solid", bd=1, selectbackground=THEME_ACTION)
listbox_peak.pack(fill="both", expand=True, pady=(0, 5))
listbox_peak.drop_target_register(DND_FILES)
listbox_peak.dnd_bind('<<Drop>>', drop_peak)

btn_frame_peak = tk.Frame(frame_peak, bg=THEME_BG)
btn_frame_peak.pack(fill="x")
tk.Button(btn_frame_peak, text="Browse", command=browse_peak, bg=THEME_BTN, fg=THEME_ACCENT, activebackground=THEME_BTN_ACTIVE, relief="flat", cursor="hand2", font=("", 9, "bold")).pack(side="left", expand=True, fill="x", padx=2)
tk.Button(btn_frame_peak, text="Remove", command=remove_peak, bg=THEME_BTN, fg=THEME_ACCENT, activebackground=THEME_BTN_ACTIVE, relief="flat", cursor="hand2", font=("", 9, "bold")).pack(side="left", expand=True, fill="x", padx=2)

# [2. MFC 셋포인트 영역 - 우측 배치]
frame_mfc = tk.LabelFrame(frame_lists, text="2. MFC 유량 제어 기록 파일 (*.csv)", bg=THEME_BG, fg=THEME_ACCENT, font=("", 10, "bold"), padx=5, pady=5)
frame_mfc.grid(row=0, column=1, sticky="nsew", padx=5)

listbox_mfc = tk.Listbox(frame_mfc, selectmode=tk.EXTENDED, height=10, relief="solid", bd=1, selectbackground=THEME_ACTION)
listbox_mfc.pack(fill="both", expand=True, pady=(0, 5))
listbox_mfc.drop_target_register(DND_FILES)
listbox_mfc.dnd_bind('<<Drop>>', drop_mfc)

btn_frame_mfc = tk.Frame(frame_mfc, bg=THEME_BG)
btn_frame_mfc.pack(fill="x")
tk.Button(btn_frame_mfc, text="Browse", command=browse_mfc, bg=THEME_BTN, fg=THEME_ACCENT, activebackground=THEME_BTN_ACTIVE, relief="flat", cursor="hand2", font=("", 9, "bold")).pack(side="left", expand=True, fill="x", padx=2)
tk.Button(btn_frame_mfc, text="Remove", command=remove_mfc, bg=THEME_BTN, fg=THEME_ACCENT, activebackground=THEME_BTN_ACTIVE, relief="flat", cursor="hand2", font=("", 9, "bold")).pack(side="left", expand=True, fill="x", padx=2)

# [초기화 버튼 영역]
frame_bot = tk.Frame(root, bg=THEME_BG)
frame_bot.pack(fill="x", padx=15)
tk.Button(frame_bot, text="Clear All Files", command=clear_all, width=15, bg="#EAECEE", fg="#5D6D7E", activebackground="#D5D8DC", relief="flat", cursor="hand2", font=("", 9, "bold")).pack(side="right", pady=5)

# [저장 설정 영역]
setting_frame = tk.LabelFrame(root, text="저장 설정", bg=THEME_BG, fg=THEME_ACCENT, font=("", 10, "bold"), padx=10, pady=10)
setting_frame.pack(pady=5, fill="x", padx=15)

use_time_format_var = tk.BooleanVar(value=True)
chk = tk.Checkbutton(setting_frame, text="파일명에 시간 자동 변환 적용 (yyMMdd_HHmm 등)", variable=use_time_format_var, bg=THEME_BG, activebackground=THEME_BG, cursor="hand2")
chk.grid(row=0, column=0, columnspan=2, sticky="w")

tk.Label(setting_frame, text="파일명 :", bg=THEME_BG).grid(row=1, column=0, sticky="w", pady=5)
entry_filename = tk.Entry(setting_frame, width=32, relief="solid", bd=1)
entry_filename.grid(row=1, column=1, sticky="w", padx=5)
tk.Label(setting_frame, text="※ 비워둘 시 기본값(yyMMdd_HHmm_가스정보) 자동 적용", fg="#7F8C8D", bg=THEME_BG, font=("", 9)).grid(row=2, column=0, columnspan=2, sticky="w", padx=5)

# [실행 버튼]
btn_process = tk.Button(root, text="Process & Save All", command=process_data, height=2, bg=THEME_ACTION, fg="white", activebackground=THEME_ACTION_ACTIVE, activeforeground="white", relief="flat", cursor="hand2", font=("", 12, "bold"))
btn_process.pack(fill="x", padx=15, pady=10)

# [하단 푸터]
footer_frame = tk.Frame(root, bg=THEME_BG)
footer_frame.pack(side="bottom", fill="x", padx=15, pady=5)

lbl_dev = tk.Label(footer_frame, text="한국공학대학교 메카트로닉스공학부 21학번 황영진", fg="#95A5A6", bg=THEME_BG, font=("", 8))
lbl_dev.pack(side="bottom", anchor="e")

lbl_github = tk.Label(footer_frame, text="GitHub Repo 🔗", fg=THEME_ACTION, bg=THEME_BG, cursor="hand2", font=("", 9, "underline"))
lbl_github.pack(side="bottom", anchor="e", pady=(0, 2))
lbl_github.bind("<Button-1>", open_github)

root.mainloop()